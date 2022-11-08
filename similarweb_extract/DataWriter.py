import csv
import logging
import os
import time
from typing import List

from openpyxl import Workbook
from slugify import slugify

logger = logging.getLogger('logger')

def fix_encoding(value):
    return value.encode('unicode_escape').decode('utf-8') if isinstance(value, str) else value

def add_time_to_filename(file_name: str) -> str:
    file_name_no_ext, ext = os.path.splitext(file_name)
    timestr = time.strftime('%Y%m%d-%H%M%S')

    suffix = ''
    n = 0
    while os.path.exists(result:=f'{file_name_no_ext}-{timestr}{suffix}{ext}'):
        n += 1
        suffix = f'-{str(n)}'

    return result

class DataWriter:
    def __init__(self, output_dir: str, file_name: str):
        self.output_dir = output_dir
        self.base_name, _ = os.path.splitext(file_name)
        self.tables = {}

    def close(self):
        pass

    def create_table(self, headers: List[str], name: str = 'results') -> None:
        pass

    def add_row(self, values: dict, table_name: str = 'results') -> None:
        self.add_rows([values], table_name)

    def add_rows(self, values: List[dict], table_name: str = 'results') -> None:
        pass

class CsvWriter(DataWriter):
    def __init__(self, output_dir: str, file_name: str) -> None:
        super().__init__(output_dir, file_name)
        self.results_folder = os.path.join(self.output_dir, f'{self.base_name}-similarweb-results')
        if not os.path.exists(self.results_folder):
            os.makedirs(self.results_folder)

    def close(self):
        for table in self.tables.values():
            file = table.get('file')
            file.close()

            file_name = table.get('file_name')
            temp_file_name = file_name + '.temp'

            file_name = add_time_to_filename(file_name)

            os.rename(temp_file_name, file_name)
            logger.info('Results stored in file: %s', file_name)


    def create_table(self, headers: List[str], name: str = 'results') -> None:
        if name in self.tables:
            raise ValueError(f'Table "{name}" already exists!')

        elif headers is not None:
            file_name = os.path.join(self.results_folder, f'{self.base_name}-{slugify(name)}.csv')
            temp_file_name = file_name + '.temp'
            try:
                os.remove(temp_file_name)
            except OSError:
                pass

            output_file = open(temp_file_name, 'w', newline='', encoding='utf-8')
            writer = csv.writer(output_file, delimiter=',', quotechar='"', quoting=csv.QUOTE_MINIMAL)
            writer.writerow(headers)

            self.tables[name] = {
                'headers': headers or [],
                'file_name': file_name,
                'file': output_file,
                'writer': writer
            }

    def add_rows(self, values: List[dict], table_name: str = 'results') -> None:
        if table_name in self.tables:
            table = self.tables.get(table_name)
            rows = [
                [value.get(header, '') for header in table.get('headers')]
                for value in values
            ]
            table.get('writer').writerows(rows)

class XlsxWriter(DataWriter):
    def __init__(self, output_dir: str, file_name: str) -> None:
        super().__init__(output_dir, file_name)
        self.results_file_path = os.path.join(self.output_dir, f'{self.base_name}-similarweb-results.xlsx')
        self.workbook = Workbook(write_only=True)

    def close(self):
        file_name = add_time_to_filename(self.results_file_path)
        self.workbook.save(file_name)
        logger.info('Results stored in file: %s', file_name)

    def create_table(self, headers: List[str], name: str = 'results') -> None:
        if name in self.tables:
            raise ValueError(f'Table id "{name}" already exists!')

        elif headers is not None:
            worksheet = self.workbook.create_sheet(name)
            worksheet.append(headers)
            self.tables[name] = {
                'headers': headers or [],
                'worksheet': worksheet,
            }

    def add_rows(self, values: List[dict], table_name: str = 'results') -> None:
        if table_name in self.tables:
            table = self.tables.get(table_name)
            worksheet = table.get('worksheet')
            for row in values:
                worksheet.append([fix_encoding(row.get(header, '')) for header in table.get('headers')])
